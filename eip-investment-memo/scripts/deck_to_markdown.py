#!/usr/bin/env python3
"""Dump a memo deck to markdown, so it can be read without opening PowerPoint.

Useful when picking up an existing memo to update or extend: you need to know what is
already in it, which sections exist, and what the numbers say, before adding anything.
Also useful for reviewing a deck you just built when no renderer is available.

    python scripts/deck_to_markdown.py memo.pptx > memo.md

Slide titles, body text, tables and chart data all come through. Position and styling do
not -- for layout problems use check_deck.py instead.
"""

import argparse
import sys

from pptx import Presentation

EMU_IN = 914400.0


def shape_sort_key(shape):
    """Reading order: top to bottom, then left to right, in half-inch bands."""
    top = (shape.top or 0) / EMU_IN
    left = (shape.left or 0) / EMU_IN
    return (round(top * 2), left)


def render_table(table):
    lines = []
    for r, row in enumerate(table.rows):
        cells = [" ".join(c.text.split()) for c in row.cells]
        lines.append("| " + " | ".join(cells) + " |")
        if r == 0:
            lines.append("|" + "|".join(["---"] * len(cells)) + "|")
    return lines


def render_chart(chart):
    lines = []
    try:
        cats = [str(c) for c in chart.plots[0].categories]
    except (IndexError, ValueError):
        cats = []
    lines.append(f"_chart ({chart.chart_type})_ categories: {', '.join(cats) or 'n/a'}")
    for series in chart.series:
        values = ", ".join("" if v is None else f"{v:g}" for v in series.values)
        lines.append(f"  - {series.name}: {values}")
    return lines


def dump(path):
    prs = Presentation(path)
    out = [f"# {path}", f"", f"{len(prs.slides)} slides", ""]

    for n, slide in enumerate(prs.slides, 1):
        out.append(f"## Slide {n}")
        out.append("")
        for shape in sorted(slide.shapes, key=shape_sort_key):
            if getattr(shape, "has_table", False):
                out += render_table(shape.table) + [""]
            elif getattr(shape, "has_chart", False):
                out += render_chart(shape.chart) + [""]
            elif shape.shape_type == 13:
                out.append(f"_[image: {shape.name}]_")
                out.append("")
            elif shape.has_text_frame:
                text = "\n".join(
                    p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()
                )
                if text:
                    out.append(text)
                    out.append("")
        out.append("")
    return "\n".join(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("-o", "--output", help="write here instead of stdout")
    args = ap.parse_args()

    text = dump(args.deck)
    if args.output:
        with open(args.output, "w") as f:
            f.write(text)
        print(f"wrote {args.output}", file=sys.stderr)
    else:
        print(text)


if __name__ == "__main__":
    main()
